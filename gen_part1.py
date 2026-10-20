"""
generate_final_report.py
Tạo báo cáo đồ án ĐẦY ĐỦ, HOÀN CHỈNH > 50 trang cho website QUẢN LÝ PHÒNG TRỌ
Giữ nguyên hình ảnh trang bìa, sửa tên đề tài, giữ nguyên Chương I, viết mới toàn diện từ Chương II.
"""
import sys
import os
sys.stdout.reconfigure(encoding='utf-8')

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC_PATH = r"c:\xampp\htdocs\QuanLyPhongTro\Mẫu báo cáo Đồ án website.docx"
OUT_PATH = r"c:\xampp\htdocs\QuanLyPhongTro\BaoCao_QuanLyPhongTro.docx"
IMG_DIR  = r"c:\xampp\htdocs\QuanLyPhongTro\assets\images"

doc = Document(SRC_PATH)

def set_font(run, name='Times New Roman', size=13, bold=False, italic=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    rPr = run._r.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:eastAsia'), name)
    rFonts.set(qn('w:ascii'), name)
    rFonts.set(qn('w:hAnsi'), name)
    rPr.insert(0, rFonts)

def add_chapter(doc, text):
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.3
    run = p.add_run(text)
    set_font(run, size=15, bold=True, color=(30, 58, 138))
    return p

def add_sec1(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.3
    run = p.add_run(text)
    set_font(run, size=13.5, bold=True, color=(30, 58, 138))
    return p

def add_sec2(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.3
    run = p.add_run(text)
    set_font(run, size=13, bold=True, color=(15, 23, 42))
    return p

def add_sec3(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.3
    run = p.add_run(text)
    set_font(run, size=12.5, bold=True, italic=True, color=(51, 65, 85))
    return p

def add_p(doc, text, bold=False, italic=False, indent=True, center=False, size=13):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent and not center:
        p.paragraph_format.first_line_indent = Cm(1.27)
    p.paragraph_format.space_after = Pt(4.5)
    p.paragraph_format.line_spacing = 1.3
    run = p.add_run(text)
    set_font(run, size=size, bold=bold, italic=italic)
    return p

def add_bullet(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(1.27)
    p.paragraph_format.first_line_indent = Cm(-0.63)
    p.paragraph_format.space_after = Pt(3.5)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run('•  ' + text)
    set_font(run, size=12.5)
    return p

def add_image(doc, filename, caption, width_in=4.8):
    path = os.path.join(IMG_DIR, filename)
    if os.path.exists(path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run()
        run.add_picture(path, width=Inches(width_in))
    else:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"[Hình ảnh: {filename}]")
        set_font(run, size=11, italic=True, color=(150, 150, 150))
    
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(8)
    r = cp.add_run(caption)
    set_font(r, size=11, bold=True, italic=True, color=(71, 85, 105))

def add_table(doc, headers, rows, caption=None):
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_before = Pt(6)
        cp.paragraph_format.space_after = Pt(4)
        r = cp.add_run(caption)
        set_font(r, size=11, bold=True, italic=True, color=(30, 58, 138))

    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = 'Table Grid'
    tbl.autofit = True

    # Header
    hr = tbl.rows[0]
    for i, h in enumerate(headers):
        c = hr.cells[i]
        c.text = ''
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(h)
        set_font(run, size=10.5, bold=True, color=(255, 255, 255))
        tcPr = c._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), '1E3A8A')
        tcPr.append(shd)

    # Rows
    for ri, row in enumerate(rows):
        tr = tbl.rows[ri + 1]
        bg_fill = 'F8FAFC' if ri % 2 == 1 else 'FFFFFF'
        for ci, val in enumerate(row):
            c = tr.cells[ci]
            c.text = ''
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ci == 0 and len(val) < 6 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(2.5)
            p.paragraph_format.space_after = Pt(2.5)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(str(val))
            set_font(run, size=10)
            if bg_fill != 'FFFFFF':
                tcPr = c._tc.get_or_add_tcPr()
                shd = OxmlElement('w:shd')
                shd.set(qn('w:val'), 'clear')
                shd.set(qn('w:color'), 'auto')
                shd.set(qn('w:fill'), bg_fill)
                tcPr.append(shd)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return tbl

def add_code(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'F1F5F9')
    pPr.append(shd)
    run = p.add_run(text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9.5)
    return p

# 1. Sửa tên đề tài trên bìa (giữ nguyên hình ảnh và khung bìa)
for p in doc.paragraphs[:30]:
    for run in p.runs:
        if 'điện thoại' in run.text.lower() or 'bán hàng' in run.text.lower() or 'đề tài:' in run.text.lower():
            run.text = '< ĐỀ TÀI: XÂY DỰNG WEBSITE QUẢN LÝ PHÒNG TRỌ >'

# 2. Cập nhật Mục lục
doc.paragraphs[27].text = 'MỤC LỤC'
set_font(doc.paragraphs[27].runs[0], size=14, bold=True)

# 3. Cập nhật Lời mở đầu
mo_dau_text = (
    "Hiện nay, cùng với sự phát triển mạnh mẽ của nền kinh tế và tốc độ đô thị hóa nhanh chóng tại các đô thị lớn, "
    "nhu cầu thuê nhà trọ, căn hộ mini và phòng trọ của sinh viên, người lao động và các gia đình trẻ ngày càng gia tăng. "
    "Tuy nhiên, đại đa số các chủ nhà trọ hiện nay vẫn đang áp dụng phương pháp quản lý truyền thống thông qua ghi chép "
    "sổ sách thủ công hoặc sử dụng các bảng tính Excel rời rạc. Phương thức quản lý thủ công này bộc lộ rất nhiều bất cập "
    "nghiêm trọng: dễ thất lạc dữ liệu, tốn nhiều thời gian và công sức ghi chỉ số điện nước hàng tháng, thường xuyên xảy ra "
    "sai sót trong quá trình tính toán cộng dồn hóa đơn, khó khăn trong việc theo dõi hạn hợp đồng thuê phòng, và đặc biệt "
    "là thiếu kênh tương tác hai chiều kịp thời giữa chủ trọ và khách thuê khi phát sinh sự cố kỹ thuật hỏng hóc.\n\n"
    "Xuất phát từ nhu cầu thực tiễn cấp thiết đó, đề tài “Xây dựng Website Quản Lý Phòng Trọ” được triển khai nhằm số hóa "
    "toàn diện quy trình vận hành khu nhà trọ. Ứng dụng được xây dựng trên nền tảng web PHP 8+ kết hợp hệ quản trị cơ sở "
    "dữ liệu MySQL và thư viện giao diện hiện đại Bootstrap 5, mang lại giải pháp quản trị tập trung, tự động hóa tính toán "
    "điện nước, lập hóa đơn, ghi nhận lịch sử thanh toán và tiếp nhận báo hỏng trực tuyến. Bản báo cáo đồ án này trình bày "
    "chi tiết từ cơ sở lý thuyết, phân tích thiết kế, cài đặt mã nguồn cho đến kiểm thử và đánh giá kết quả thực nghiệm."
)
if len(doc.paragraphs) > 54:
    doc.paragraphs[54].text = mo_dau_text
    set_font(doc.paragraphs[54].runs[0], size=12.5)

# 4. Xóa từ Chương II cũ trở đi
to_del = []
found_ch2 = False
for p in doc.paragraphs:
    if not found_ch2 and ('Chương II' in p.text or 'chương II' in p.text.lower()):
        found_ch2 = True
    if found_ch2:
        to_del.append(p)

for p in to_del:
    p._element.getparent().remove(p._element)

print("Đã làm sạch Chương II cũ, bắt đầu viết nội dung chi tiết...")
