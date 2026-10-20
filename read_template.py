import sys
sys.stdout.reconfigure(encoding='utf-8')
import docx

doc = docx.Document(r'c:\xampp\htdocs\QuanLyPhongTro\template.docx')
print(f'Sections: {len(doc.sections)}')
print(f'Paragraphs: {len(doc.paragraphs)}')
print(f'Tables: {len(doc.tables)}')
print('---ALL paragraphs (non-empty)---')
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(f'[{i}] Style={p.style.name!r} | {p.text[:150]!r}')
